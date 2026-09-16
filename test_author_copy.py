"""Two files from one copyedit: the team's and the author's.

The quality team's request, 16 Sep 2026 — they can read a redline carrying every
guard note and every house-style point; an author opening the same file cannot.
"""
import docx
import pytest
from docx.oxml.ns import qn

from editor import AUTHOR, generate_redline_docx


@pytest.fixture
def manuscript(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("A Study of Something")
    doc.add_paragraph("The sample was heated to 45 degrees and left for an hour.")
    doc.add_paragraph("Figure 1. Plot for T4 (80g) of the dried stem.")
    path = str(tmp_path / "m.docx")
    doc.save(path)
    return path


QUERIES = [
    {"index": 2, "query": "The copyedit changed `T4` in this caption.",
     "suggestion": None, "audience": "author"},
    {"index": 1, "query": "The copyedit removed the day and month; restored.",
     "suggestion": None, "audience": "internal"},
    {"index": 1, "query": "A house-style note with no audience set at all.",
     "suggestion": None},
]


def _comments(path):
    with docx.Document(path).part.package.parts[0].blob and open(path, "rb"):
        pass
    import zipfile
    with zipfile.ZipFile(path) as z:
        if "word/comments.xml" not in z.namelist():
            return ""
        return z.read("word/comments.xml").decode("utf-8")


def _edited(paragraphs):
    out = list(paragraphs)
    return out


def test_the_internal_copy_carries_every_query(manuscript, tmp_path):
    out = str(tmp_path / "internal.docx")
    generate_redline_docx(manuscript, ["A Study of Something",
                                       "The sample was heated to 45 °C and left for "
                                       "an hour.",
                                       "Figure 1. Plot for T4 (80 g) of the dried stem."],
                          out, queries=QUERIES)
    body = _comments(out)
    assert "caption" in body and "day and month" in body and "house-style" in body


def test_the_author_copy_carries_only_their_questions(manuscript, tmp_path):
    out = str(tmp_path / "author.docx")
    generate_redline_docx(manuscript, ["A Study of Something",
                                       "The sample was heated to 45 °C and left for "
                                       "an hour.",
                                       "Figure 1. Plot for T4 (80 g) of the dried stem."],
                          out, queries=QUERIES, audience=AUTHOR)
    body = _comments(out)
    assert "caption" in body, "the author's own question must be there"
    assert "day and month" not in body, "a guard note is the team's business"
    assert "house-style" not in body, "an untagged query is not an author's question"


def test_the_author_copy_highlights_what_changed(manuscript, tmp_path):
    """A manuscript read with markup off looks untouched. The decision is still the
    author's — every change is a tracked change they accept or reject in Word."""
    out = str(tmp_path / "author.docx")
    generate_redline_docx(manuscript, ["A Study of Something",
                                       "The sample was heated to 45 °C and left for "
                                       "an hour.",
                                       "Figure 1. Plot for T4 (80 g) of the dried stem."],
                          out, queries=QUERIES, audience=AUTHOR)
    doc = docx.Document(out)
    inserted = list(doc.element.body.iter(qn("w:ins")))
    assert inserted, "the fixture must produce at least one tracked insertion"
    highlighted = [h for h in doc.element.body.iter(qn("w:highlight"))]
    assert highlighted, "inserted text is not highlighted"


def test_the_internal_copy_is_not_highlighted(manuscript, tmp_path):
    out = str(tmp_path / "internal.docx")
    generate_redline_docx(manuscript, ["A Study of Something",
                                       "The sample was heated to 45 °C and left for "
                                       "an hour.",
                                       "Figure 1. Plot for T4 (80 g) of the dried stem."],
                          out, queries=QUERIES)
    doc = docx.Document(out)
    assert not list(doc.element.body.iter(qn("w:highlight")))


def test_both_copies_carry_the_same_tracked_changes(manuscript, tmp_path):
    """One copyedit, two readerships — never two different manuscripts."""
    edited = ["A Study of Something",
              "The sample was heated to 45 °C and left for an hour.",
              "Figure 1. Plot for T4 (80 g) of the dried stem."]
    a, b = str(tmp_path / "i.docx"), str(tmp_path / "a.docx")
    generate_redline_docx(manuscript, edited, a, queries=QUERIES)
    generate_redline_docx(manuscript, edited, b, queries=QUERIES, audience=AUTHOR)
    text = lambda p: [q.text for q in docx.Document(p).paragraphs]
    assert text(a) == text(b)
