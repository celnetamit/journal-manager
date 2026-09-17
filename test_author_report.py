"""The one page an author gets, and what must and must not be on it."""

from __future__ import annotations

import author_report as A


class Finding:
    def __init__(self, rule, message=""):
        self.rule, self.message = rule, message


_REVIEW = """# AI Peer Review Report
## Summary
The paper studies bioremediation of crude-oil-contaminated soil.
## Significance & Originality
The topic is well covered in the literature.
## Strengths
- A clear experimental design.
## Major Concerns
- The control group is not described.
## Minor Issues
- Figure 2 is illegible at the printed size.
## Recommendation
Major Revision — the controls must be described.
## Confidence
Medium.
"""


def test_the_missing_things_are_named_and_the_formatting_ones_are_not():
    report = A.build(
        "paper.docx",
        findings=[Finding("front.abstract-missing"),
                  Finding("table.cited-but-missing",
                          "the text refers to Table 6, and no such caption is in "
                          "the manuscript"),
                  Finding("references.hanging-indent", "0.25\" hanging indent")],
        ai_review_md=_REVIEW)
    assert "The Abstract is missing." in report
    assert "Table 6" in report
    # A hanging indent is the production team's business, not the author's.
    assert "hanging indent" not in report


def test_a_cited_work_with_no_reference_reaches_the_author():
    report = A.build("paper.docx", queries=[{
        "guard": "refuse_citations_without_a_reference",
        "query": "`FAO (2015)` is cited here but there is no reference for it in the "
                 "list. The author's own citation has been kept as it was.",
    }])
    assert "FAO (2015)" in report and "no reference" in report


def test_the_review_is_cut_to_what_the_author_acts_on():
    report = A.build("paper.docx", ai_review_md=_REVIEW)
    assert "The control group is not described." in report
    assert "Major Revision" in report
    # The referee's own confidence is for whoever decides the paper.
    assert "Confidence" not in report
    assert "Medium." not in report


def test_a_clean_manuscript_says_so_rather_than_showing_an_empty_heading():
    report = A.build("paper.docx", findings=[], queries=[], ai_review_md=_REVIEW)
    assert "Nothing." in report


def test_the_same_gap_reported_twice_is_written_once():
    finding = Finding("front.abstract-missing")
    report = A.build("paper.docx", findings=[finding, finding])
    assert report.count("The Abstract is missing.") == 1


def test_no_review_says_so_instead_of_leaving_a_blank_section():
    report = A.build("paper.docx", ai_review_md="")
    assert "No review was produced" in report


def test_the_report_always_points_at_the_redline():
    """A report that does not say where the edits are sends the author hunting."""
    report = A.build("paper.docx")
    assert "tracked-changes" in report
    # And it names the signature the comments actually carry, so they can be found.
    assert "Query to Author" in report


def test_one_fault_repeated_is_one_line():
    """Six uncaptioned tables produced six identical lines, which buried the line
    saying the Abstract was not there."""
    report = A.build("paper.docx", findings=[
        Finding("table.caption", f"table {n} has no 'Table N' caption immediately "
                                 f"above it") for n in range(1, 7)
    ] + [Finding("front.abstract-missing")])
    assert "6 tables have no 'Table N' caption above them." in report
    assert report.count("caption above them") == 1
    assert "The Abstract is missing." in report


def test_the_parts_of_the_paper_come_before_the_artwork():
    report = A.build("paper.docx", findings=[
        Finding("table.cited-but-missing", "the text refers to Table 6"),
        Finding("front.abstract-missing"),
    ])
    assert report.index("Abstract is missing") < report.index("Table 6")


def test_every_kind_of_bullet_renders_as_a_bullet(tmp_path):
    """Job #112's author report printed a literal `*` in front of every Minor Issue:
    the review model writes `*   ` and only `- ` was being rendered."""
    import docx
    from editor import markdown_to_docx
    out = tmp_path / "r.docx"
    markdown_to_docx("# T\n\n- dash item\n*   star item\n+ plus item\n"
                     "\nThe symbol `h(t)` is estimated.\n", str(out))
    d = docx.Document(str(out))
    bullets = [p.text for p in d.paragraphs if p.style.name == "List Bullet"]
    assert bullets == ["dash item", "star item", "plus item"]
    body = " ".join(p.text for p in d.paragraphs)
    assert "*" not in body and "`" not in body
    assert "h(t)" in body
