"""The ORCID iD section that sits just above the references.

Asked for by the quality team on 16 Sep 2026, with a typeset page as the specification:

    ORCID ID
    Md. Al-Mamun:  [iD] https://orcid.org/0000-0002-9970-3122
    Md. Mashiur Rahman:  [iD] https://orcid.org/0009-0008-5205-279X

Authors put their iD in the front matter — usually in an author-details block — and it
then sits in the middle of the paper doing nothing. Collected into one place above the
references, with the green mark linking to the record, it is what a reader and an indexer
both look for.

Two rules this file keeps:

* **Nothing is invented.** An iD is only shown if the author wrote it, and it is shown
  against the name the author attached it to. A checksum is verified before anything is
  printed — ORCID's own mod-11-2 — so a mistyped digit is left out rather than published
  as somebody else's record.
* **The paragraph list is only ever added to, at the end of the build.** Tracked changes
  are aligned positionally; inserting paragraphs before that alignment would move every
  change onto the wrong text. This runs on the finished document, immediately before it
  is saved.
"""

from __future__ import annotations

import os
import re
from typing import Dict, List, Optional, Tuple

from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt

#: `0000-0002-9970-3122`, allowing the X checksum and any spacing the author typed.
_ID = re.compile(r"(\d{4})\s*-\s*(\d{4})\s*-\s*(\d{4})\s*-\s*(\d{3}[\dXx])")

#: "ORCID", "ORCID iD", "ORCID ID:" — the label an author writes before their iD.
_LABEL = re.compile(r"(?i)\borcid\b")

#: `Author 1: Jane Doe`, `Corresponding author: Jane Doe`, or a bare name line.
_NAMED = re.compile(
    r"(?i)^\s*(?:\d+\s*[.)]\s*)?(?:corresponding\s+author|author\s*\d*|name)\s*[:\-–]\s*(.+)$")

#: Where the references begin — the section this one goes above.
_REFERENCES = re.compile(r"(?i)^\s*(?:\d+[.)]?\s*)?(references|bibliography|works cited)\s*:?\s*$")

ICON = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "orcid-id.png")


def checksum_ok(digits: str) -> bool:
    """ORCID's ISO 7064 mod 11-2 check digit.

    Verified rather than trusted: a transposed pair in a sixteen-digit number is
    invisible to a reader and points at a different person's record. Publishing that
    under somebody's name is worse than leaving the iD out.
    """
    body = digits.replace("-", "")
    if len(body) != 16:
        return False
    total = 0
    for ch in body[:15]:
        if not ch.isdigit():
            return False
        total = (total + int(ch)) * 2
    remainder = total % 11
    expected = (12 - remainder) % 11
    return body[15].upper() == ("X" if expected == 10 else str(expected))


def _clean_name(raw: str) -> str:
    name = re.sub(r"\s+", " ", raw or "").strip(" .,:;–-")
    # Affiliation digits and asterisks belong to the byline, not to the name.
    name = re.sub(r"[\d*†‡§¶]+$", "", name).strip()
    return name


def find_orcids(paragraphs: List[str]) -> List[Tuple[str, str]]:
    """`[(author name, iD)]` in the order the author listed them.

    The name is the one nearest above the iD — in these manuscripts the block reads
    `Author 2: Md. Mashiur Rahman / Affiliation: … / ORCID: https://orcid.org/…`. When
    the iD sits on the same line as a name, that name wins.
    """
    found: List[Tuple[str, str]] = []
    seen: Dict[str, str] = {}
    last_name = ""

    for raw in paragraphs:
        line = (raw or "").strip()
        if not line:
            continue

        named = _NAMED.match(line)
        if named:
            candidate = _clean_name(_ID.sub("", _LABEL.sub("", named.group(1))))
            if candidate:
                last_name = candidate

        match = _ID.search(line)
        if not match:
            # A short line of capitalised words, with no label, is a name in a byline.
            if not named and len(line) < 80 and re.match(r"^[A-Z][A-Za-z.\-' ]+$", line):
                last_name = _clean_name(line)
            continue

        digits = "-".join(match.groups()).upper()
        if not checksum_ok(digits):
            continue

        # `Jane Doe: https://orcid.org/0000-…` — the name is on this line.
        before = line[:match.start()]
        inline = _clean_name(_LABEL.sub("", before).strip(" :–-"))
        inline = re.sub(r"https?://\S*$", "", inline).strip(" :–-")
        name = inline if 2 < len(inline) < 80 else last_name
        if not name:
            continue
        if seen.get(name) == digits or digits in seen.values():
            continue
        seen[name] = digits
        found.append((name, digits))
    return found


def _hyperlink(paragraph, url: str):
    """An empty `w:hyperlink` in this paragraph, ready to be given runs."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    paragraph._p.append(link)
    return link


def _linked_run(paragraph, url: str, *, text: str = "", image: Optional[str] = None):
    """A run inside a hyperlink — an icon, a label, or both."""
    link = _hyperlink(paragraph, url)
    run = paragraph.add_run()
    if image:
        run.add_picture(image, height=Pt(11))
    if text:
        run.text = text
        run.font.color.rgb = None
        run.font.underline = True
    link.append(run._r)
    return run


def add_orcid_section(document, pairs: List[Tuple[str, str]]) -> int:
    """Put the ORCID block immediately above the references. Returns how many iDs.

    Placed above the references because that is where this house puts it, and because it
    is the last thing about the authors before the paper stops being theirs and starts
    being other people's work.
    """
    if not pairs:
        return 0

    # Already there? Then the author (or a previous run) has written it, and a second
    # copy is worse than none.
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if _LABEL.match(text) and len(text) < 20:
            return 0

    anchor = None
    for para in document.paragraphs:
        if _REFERENCES.match((para.text or "").strip()):
            anchor = para
            break

    def new_paragraph():
        para = document.add_paragraph()
        if anchor is not None:
            anchor._p.addprevious(para._p)
        return para

    heading = new_paragraph()
    run = heading.add_run("ORCID ID")
    run.bold = True

    for name, digits in pairs:
        url = f"https://orcid.org/{digits}"
        line = new_paragraph()
        line.add_run(f"{name}: ")
        _linked_run(line, url, image=ICON if os.path.exists(ICON) else None)
        line.add_run(" ")
        _linked_run(line, url, text=url)
    return len(pairs)
