"""An inline equation survives the copyedit — job #100.

Built as a real `.docx` with real OMML rather than a mock, because every part of this
failed at the XML level: `Paragraph.text` steps over an equation, `Paragraph.clear()`
deletes it, and neither of those is visible from a stub.
"""
import re

import docx
import pytest

import edit_guards as G
from editor import (OBJECT_PLACEHOLDER, generate_redline_docx,
                    paragraph_objects, paragraph_text_with_objects, read_docx)

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"



W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _redline_text(p) -> str:
    """The paragraph as it reads on the page, tracked changes and equations included.

    `paragraph_text_with_objects` walks the direct `w:r` children, which is right for
    the manuscript and wrong for the redline: there the edited words live inside
    `w:ins` and `w:del`, one level down.
    """
    out = []
    for node in p._p.iter():
        if node.tag == f"{{{W}}}t":
            out.append(node.text or "")
        elif node.tag == f"{{{M}}}oMath":
            out.append("[EQ]")
    return "".join(out)


def _add_equation(paragraph, text):
    """An inline OMML equation, as Word writes one."""
    from docx.oxml import parse_xml
    xml = (f'<m:oMath xmlns:m="{M}" '
           f'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f'<m:r><m:t>{text}</m:t></m:r></m:oMath>')
    paragraph._p.append(parse_xml(xml))


def _manuscript(tmp_path):
    """`where [K_IC ≈ 0.7 MPa] is the fracture toughness of silica, [a = 50 nm] is
    the characteristic flaw size` — job #100's sentence, in miniature."""
    doc = docx.Document()
    doc.add_paragraph("Title of the Paper")
    doc.add_paragraph("Abstract: a short abstract.")
    p = doc.add_paragraph()
    p.add_run("where ")
    _add_equation(p, "K_IC ≈ 0.7 MPa")
    p.add_run(" is the fracture toughness of silica, and ")
    _add_equation(p, "a = 50 nm")
    p.add_run(" is the characteristic flaw size.")
    path = str(tmp_path / "manuscript.docx")
    doc.save(path)
    return path


def test_the_copyedit_is_shown_the_equation_not_a_gap(tmp_path):
    paras = read_docx(_manuscript(tmp_path))
    sentence = [t for t in paras if "fracture toughness" in t][0]
    assert sentence.count(OBJECT_PLACEHOLDER) == 2
    # And no run of blanks where an equation was: that hole is what the model filled
    # with `K_Ic` and `a` on the real job.
    assert not re.search(r"\s{2,}", sentence)


def test_the_equation_survives_a_paragraph_the_copyedit_rewrote(tmp_path):
    path = _manuscript(tmp_path)
    original = read_docx(path)
    edited = [t.replace("where", "Where").replace("silica,", "silica")
              for t in original]
    out = str(tmp_path / "redline.docx")
    generate_redline_docx(path, edited, out)

    doc = docx.Document(out)
    sentence = [p for p in doc.paragraphs
                if "fracture toughness" in paragraph_text_with_objects(p)][0]
    maths = paragraph_objects(sentence)
    assert len(maths) == 2, "the author's equations were dropped by the rewrite"
    assert "0.7 MPa" in "".join(t.text or "" for t in maths[0].iter(f"{{{M}}}t"))
    assert "50 nm" in "".join(t.text or "" for t in maths[1].iter(f"{{{M}}}t"))
    # The placeholder is a working character, not something to ship to the author.
    assert OBJECT_PLACEHOLDER not in doc.element.xml


def test_the_equation_stays_in_the_middle_of_its_sentence(tmp_path):
    path = _manuscript(tmp_path)
    original = read_docx(path)
    edited = [t.replace("where", "Where") for t in original]
    out = str(tmp_path / "redline.docx")
    generate_redline_docx(path, edited, out)

    doc = docx.Document(out)
    p = [p for p in doc.paragraphs if "fracture toughness" in p.text][0]
    rendered = _redline_text(p)
    before, after = rendered.split("[EQ]")[0], rendered.split("[EQ]")[-1]
    assert "here" in before, rendered
    assert "characteristic flaw size" in after, rendered


def test_a_copyedit_that_loses_an_equation_is_refused():
    """The model's own job #100 output: the equations replaced by its guesses."""
    original = ["where ￼ is the fracture toughness of silica, and ￼ is the flaw size"]
    edited = ["where K_Ic is the fracture toughness of silica, and a is the flaw size"]
    out, queries = G.keep_every_equation(original, edited)
    assert out == original
    assert len(queries) == 1 and "equation" in queries[0]["query"]


def test_an_ordinary_edit_beside_an_equation_is_kept():
    original = ["where ￼ is the fracture toughness of silica"]
    edited = ["Where ￼ is the fracture toughness of silica"]
    out, queries = G.keep_every_equation(original, edited)
    assert out == edited and queries == []


def test_the_placeholder_never_reaches_a_reader(tmp_path):
    """It is a working character between the reader and the writer. The JATS goes to
    a typesetter and an indexer, where a stray U+FFFC is a defect in the file."""
    from editor import build_jats_xml, for_display
    xml = build_jats_xml(["A Title", "where ￼ is the toughness of silica."])
    assert OBJECT_PLACEHOLDER not in xml
    assert "[equation]" in xml
    assert for_display("a ￼ b") == "a [equation] b"
