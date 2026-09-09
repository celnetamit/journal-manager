"""What job #53 lost, and the guards that now stop it losing it.

Every string here is taken from *Hybrid Braking System.docx* and its redline — the
originals reconstructed from the tracked deletions, the results from what the redline
actually rendered. Three of these were not "the copyedit made a poor choice": they were
the author's content leaving the document with nothing saying so.
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from docx import Document
from docx.shared import Inches

import edit_guards
import proofread as P

# The smallest valid PNG, so a test can build a real inline image.
_PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
        b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00"
        b"\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")


def _drawings(doc) -> int:
    buf = io.BytesIO()
    doc.save(buf)
    xml = zipfile.ZipFile(io.BytesIO(buf.getvalue())).read("word/document.xml").decode()
    return len(re.findall(r"<w:drawing[ >]", xml))


# ---------------------------------------------------------------- figures survive


@pytest.mark.parametrize("image_first", [False, True])
def test_copyediting_a_caption_does_not_delete_the_figure(image_first):
    """The bug that cost job #53 three of its four figures.

    The caption and the picture shared a paragraph, `_mark_up_paragraph` cleared the
    paragraph to rewrite the text, and `Paragraph.clear()` takes the `w:drawing` with it.
    Parameterised over both orders because a figure sits either side of its caption and
    only one of those was ever likely to be tried by hand.
    """
    import editor

    doc = Document()
    p = doc.add_paragraph()
    if image_first:
        p.add_run().add_picture(io.BytesIO(_PNG), width=Inches(1))
    p.add_run("Fig 1: actual model of hybrid breaking system")
    if not image_first:
        p.add_run().add_picture(io.BytesIO(_PNG), width=Inches(1))

    assert _drawings(doc) == 1
    editor._mark_up_paragraph(
        doc.paragraphs[0],
        "Fig 1: actual model of hybrid breaking system",
        "Figure 1. Actual model of hybrid braking system.",
        1)
    assert _drawings(doc) == 1, "the figure must survive its caption being copyedited"


def test_the_copyedit_still_happens_around_the_figure():
    """Protecting the picture must not cost the edit — both have to land."""
    import editor

    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Fig 3: Proportional-Integral-Derivative")
    p.add_run().add_picture(io.BytesIO(_PNG), width=Inches(1))
    editor._mark_up_paragraph(doc.paragraphs[0],
                              "Fig 3: Proportional-Integral-Derivative",
                              "Figure 3. Proportional-integral-derivative.", 1)
    buf = io.BytesIO()
    doc.save(buf)
    xml = zipfile.ZipFile(io.BytesIO(buf.getvalue())).read("word/document.xml").decode()
    assert "<w:ins " in xml and "<w:del " in xml     # tracked, not silently replaced
    assert _drawings(doc) == 1


# ------------------------------------------------------------- caption labels stay


def test_the_table_label_is_not_stripped_as_heading_numbering():
    """`Table 2: Comparative Summary Table` -> `Comparative Summary Table`.

    The house heading rule removes leading numbering, and "Table 2:" looks exactly like
    it. The manuscript then failed its own cross-reference check for a caption the tool
    had itself deleted.

    A caption stays a caption and keeps its corrections: only the label is put back, so
    the spelling fixes inside the caption survive alongside it. Reverting the whole
    paragraph would have thrown those away together with the damage.
    """
    out, queries = edit_guards.restore_protected_text(
        ["Table 2: Comparision of breaking sytem"], ["Comparison of braking system"])
    assert out == ["Table 2. Comparison of braking system"]
    assert queries and "Table 2" in queries[0]["query"]


def test_the_table_label_is_restored_without_losing_the_copyedit():
    """The label alone comes back; the edited wording stays."""
    out, _q = edit_guards.restore_protected_text(
        ["Table 2: Comparative Summary Table"], ["Comparative Summary Table"])
    assert out == ["Table 2. Comparative Summary Table"]


def test_a_caption_sharing_a_paragraph_with_a_heading_is_not_dropped():
    """Job #53 ¶36: the heading survived, the figure caption did not."""
    before = "3.2 Sensor Fusion                        Fig: 2 Design of Prototype"
    out, queries = edit_guards.restore_protected_text([before], ["Sensor Fusion"])
    assert out == [before]
    assert "Figure 2" in queries[0]["query"]
    assert "split them into two paragraphs" in queries[0]["query"]


def test_a_caption_may_still_be_reformatted():
    """`Fig 1: actual model` -> `Figure 1. Actual model.` is the house rule working.

    The guard asks only whether the element is still called Figure 1 — if it flagged
    reformatting it would fire on every caption in every manuscript.
    """
    out, queries = edit_guards.restore_protected_text(
        ["Fig 1: actual model of hybrid breaking system"],
        ["Figure 1. Actual model of hybrid braking system."])
    assert out == ["Figure 1. Actual model of hybrid braking system."]
    assert queries == []


def test_prose_mentioning_a_table_is_not_treated_as_a_caption():
    """Body text is edited freely; only a lost *number* matters, and none is lost here."""
    out, queries = edit_guards.restore_protected_text(
        ["Table 1 illustrate the comparision of breaking system."],
        ["Table 1 illustrates the comparison of braking systems."])
    assert queries == []


# ------------------------------------------------------------ duplicate references


def test_the_same_book_listed_twice_is_caught():
    """Entries 6 and 8 of job #53 are the same work; each is well-formed alone.

    These two strings are verbatim from the job's own output, and they carry no volume,
    issue or page numbers — which matters, because a duplicate check written against
    tidy `35(1): 55-65` citations would pass its tests and still have missed this pair.
    The proofreader runs on the *edited* paragraphs, which is where the reference rules
    have already put the `6.` / `8.` markers on, so that is what is fed in here.
    """
    refs = [
        "5. Meng B, et al. A survey of brake-by-wire system. IEEE Access. 2020; 8(2): 22–35p.",
        "6. Aggarwal G. Integrated Technologies in Electrical, Electronics and "
        "Biotechnology Engineering. 2025.",
        "7. Topping KJ. Effectiveness of blended learning. Rev Educ. 2022; 10(2): 3–19p.",
        "8. Aggarwal G. Integrated Technologies in Electrical, Electronics and "
        "Biotechnology Engineering. 2025.",
    ]
    hits = [f for f in P._reference_duplicate_findings(refs)
            if f.rule == "reference.duplicate"]
    assert len(hits) == 1, "reported once, against the second copy"
    assert hits[0].paragraph == 3
    assert "paragraph 2" in hits[0].message          # points back at the first

def test_two_different_works_by_one_author_are_not_duplicates():
    """The commonest way a duplicate check earns its own ignoring."""
    refs = [
        "6. Aggarwal G. Integrated Technologies in Electrical Engineering. "
        "J Elec Syst. 2025; 4(1): 11–20p.",
        "8. Aggarwal G. Advances in Biotechnology Instrumentation. "
        "J Biotech Methods. 2024; 3(2): 45–58p.",
    ]
    assert P._reference_duplicate_findings(refs) == []


def test_a_duplicate_is_found_through_differing_punctuation():
    """The same entry retyped is rarely retyped character-for-character."""
    refs = [
        "[3] Zhang Y., Zhao C., & Li Z. Electric vehicle regenerative braking system "
        "simulation. Chin Autom Congr. 2020; 2(1): 5–14p.",
        "[9]. Zhang Y, Zhao C, Li Z. Electric vehicle regenerative braking system "
        "simulation. Chin Autom Congr. 2020; 2(1): 5-14p.",
    ]
    assert len(P._reference_duplicate_findings(refs)) == 1


# ------------------------------------------------------------------ DOIs in prose


def test_a_doi_is_only_offered_to_a_bibliography_entry():
    """Job #53 got Crossref DOIs appended to its literature review.

    The old test was "over 30 characters and contains a year", which is most of any
    manuscript's discussion section.
    """
    import editor

    prose = ('Xing, C., et al. (2024). "Regenerative Braking Control Strategy of '
             'Vehicle With In-Wheel Motor Drive System" proposed a fault-tolerant '
             'approach that this study builds on.')
    entry = ("4. Mikropoulos TA, Iatraki G. Digital technology supports science "
             "education. Educ Inf Technol. 2023; 28(4): 3911–3935p.")
    assert not editor._looks_like_reference_entry(prose)
    assert editor._looks_like_reference_entry(entry)


# ------------------------------------------- spelling survives a rejected copyedit

import losscheck

# Job #53 ¶68, verbatim: a heading, a source note, a citation and a body sentence, all
# typed into one paragraph. The copyedit shortened it by 37%, so the loss guard rejected
# the whole thing — and the two spelling fixes inside it went back with the rest.
_P68 = ("6.1 Theoretical Performance Metrics  (Derived from theoretical model and "
        "prototype testing as discussed in Zhang et al., 2022; Guo et al., 2021) [2, 5]"
        "Table 1 illustrate the comparision of breaking system between conventional "
        "disc brake, electromagnetic brake, regenerative brake and hybrid braking system.")
_P68_EDITED = ("Theoretical Performance Metrics [2, 5] Table 1 illustrates the comparison "
               "of braking system between conventional disc brake, electromagnetic "
               "brake, regenerative brake and hybrid braking system.")


def test_the_loss_guard_still_rejects_this_edit():
    """The premise: without this the paragraph would never have been reverted at all."""
    assert losscheck.check_paragraph(_P68, _P68_EDITED)["kind"] == "truncated"


def test_a_rejected_copyedit_keeps_its_spelling_corrections():
    kept = losscheck.salvage_safe_corrections(_P68, _P68_EDITED)
    assert "comparision" not in kept and "comparison" in kept
    assert "breaking system" not in kept and "braking system" in kept
    assert "illustrates" in kept                      # illustrate -> illustrates


def test_the_salvage_gives_nothing_away():
    """Everything the copyedit tried to delete is still there."""
    kept = losscheck.salvage_safe_corrections(_P68, _P68_EDITED)
    assert "Derived from theoretical model" in kept   # the source note it dropped
    assert kept.startswith("6.1")                     # the heading number it dropped
    assert "Zhang et al., 2022" in kept
    assert len(kept) >= len(_P68) - 5                 # a word swap, not a trim


def test_a_word_swapped_for_a_different_word_is_refused():
    """`disc` -> `drum` is a change of meaning, and this is not the place to take it."""
    kept = losscheck.salvage_safe_corrections(
        "The conventional disc brake was used throughout the trials.",
        "The drum brake was used.")
    assert "disc" in kept and "drum" not in kept


def test_numbers_are_never_treated_as_spelling():
    kept = losscheck.salvage_safe_corrections(
        "The rotor reached 1450 rpm during the third run of the test.",
        "The rotor reached 1650 rpm.")
    assert "1450" in kept and "1650" not in kept
