"""The ORCID block above the references.

Asked for by the quality team on 16 Sep 2026, with a typeset page as the specification:
a heading, then one line per author — name, the green iD mark, and the full
`https://orcid.org/…` — all of it linking to the record.

The rule worth testing hardest is the one nobody sees: an iD is only shown if the author
wrote it *and* it passes ORCID's own checksum. A transposed pair of digits reads as a
perfectly ordinary iD and points at a different person's record, under this author's name.
"""

import pytest

from orcid import add_orcid_section, checksum_ok, find_orcids


# --- the checksum ------------------------------------------------------------------


def test_real_ids_pass():
    assert checksum_ok("0000-0002-9970-3122")
    assert checksum_ok("0009-0008-5205-279X"), "the X check digit is a real one"
    assert checksum_ok("0009-0007-3030-1508")


def test_a_transposed_digit_is_refused():
    """The failure this exists for: it looks exactly like an iD and is somebody else's."""
    assert not checksum_ok("0000-0002-9907-3122")
    assert not checksum_ok("0000-0002-9970-3123")


# --- finding them ------------------------------------------------------------------


def test_an_author_details_block():
    paragraphs = [
        "Author 1: Md. Al-Mamun",
        "Affiliation: Bangladesh Jute Research Institute",
        "ORCID: https://orcid.org/0000-0002-9970-3122",
        "Author 2: Md. Mashiur Rahman",
        "Affiliation: University for Development Studies",
        "ORCID: https://orcid.org/0009-0008-5205-279X",
    ]
    assert find_orcids(paragraphs) == [
        ("Md. Al-Mamun", "0000-0002-9970-3122"),
        ("Md. Mashiur Rahman", "0009-0008-5205-279X"),
    ]


def test_a_name_on_the_same_line():
    assert find_orcids(["Md. Al-Mamun: https://orcid.org/0000-0002-9970-3122"]) == [
        ("Md. Al-Mamun", "0000-0002-9970-3122")]


def test_an_id_with_no_name_anywhere_is_left_out():
    """Printed under no name it is nobody's, and printed under the wrong one it is worse."""
    assert find_orcids(["ORCID: https://orcid.org/0000-0002-9970-3122"]) == []


def test_a_mistyped_id_is_left_out_rather_than_published():
    paragraphs = ["Author 1: Md. Al-Mamun", "ORCID: 0000-0002-9907-3122"]
    assert find_orcids(paragraphs) == []


def test_the_same_id_is_not_listed_twice():
    paragraphs = [
        "Author 1: Md. Al-Mamun",
        "ORCID: https://orcid.org/0000-0002-9970-3122",
        "Corresponding author: Md. Al-Mamun",
        "ORCID: https://orcid.org/0000-0002-9970-3122",
    ]
    assert len(find_orcids(paragraphs)) == 1


# --- the block itself --------------------------------------------------------------


@pytest.fixture
def document(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("CONCLUSION")
    doc.add_paragraph("Jute has evolved from a conventional packaging fibre.")
    doc.add_paragraph("REFERENCES")
    doc.add_paragraph("1. Mohanty AK, Misra M. Natural fibers. 2005.")
    return doc


def test_the_block_goes_above_the_references(document):
    added = add_orcid_section(document, [("Md. Al-Mamun", "0000-0002-9970-3122")])
    texts = [p.text for p in document.paragraphs]

    assert added == 1
    assert texts.index("ORCID ID") < texts.index("REFERENCES")
    assert texts.index("CONCLUSION") < texts.index("ORCID ID")
    assert any("Md. Al-Mamun" in t and "0000-0002-9970-3122" in t for t in texts)


def test_every_id_is_a_link(document):
    add_orcid_section(document, [("Md. Al-Mamun", "0000-0002-9970-3122")])
    xml = document.element.xml

    assert xml.count("hyperlink") >= 2, "the mark and the address both carry the link"
    assert "0000-0002-9970-3122" in xml


def test_the_green_mark_is_there(document):
    add_orcid_section(document, [("Md. Al-Mamun", "0000-0002-9970-3122")])
    assert "graphicData" in document.element.xml, "the iD mark is an image, not a letter"


def test_nothing_is_added_when_no_author_gave_one(document):
    before = len(document.paragraphs)
    assert add_orcid_section(document, []) == 0
    assert len(document.paragraphs) == before


def test_a_paper_that_already_lists_them_is_left_alone(document):
    """Typeset papers come back through this. Two ORCID blocks is worse than none."""
    para = document.add_paragraph("ORCID ID")
    document.paragraphs[2]._p.addprevious(para._p)

    assert add_orcid_section(document, [("Md. Al-Mamun", "0000-0002-9970-3122")]) == 0
