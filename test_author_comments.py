"""Whose question is it? On the author's copy, Word has to say so.

The quality team, 17 Sep 2026: the author's redline carries only the author's queries,
and every one of them was still signed `AI Editor`. To an author that reads as somebody
else's note they have been copied on, which is exactly the wrong signal for the one set
of questions only they can answer.
"""

from __future__ import annotations

import zipfile

import docx
import editor


def _redline(tmp_path, audience, text="The sample was heated to 40 and held."):
    source = tmp_path / "in.docx"
    d = docx.Document()
    d.add_paragraph(text)
    d.save(str(source))
    out = tmp_path / f"out-{audience or 'internal'}.docx"
    editor.generate_redline_docx(
        str(source), [text], str(out),
        queries=[{"index": 0, "query": "40 what — °C?", "guard": "x",
                  "audience": "author", "suggestion": None}],
        audience=audience)
    with zipfile.ZipFile(str(out)) as z:
        return z.read("word/comments.xml").decode("utf8")


def test_the_authors_copy_names_the_author(tmp_path):
    xml = _redline(tmp_path, "author")
    assert "Query to Author" in xml
    assert "AI Editor" not in xml
    # The label and the note are separate runs now — the label is bold — so the two
    # are asserted separately rather than as one string.
    assert "Author:" in xml and "40 what" in xml


def test_the_internal_copy_is_unchanged(tmp_path):
    xml = _redline(tmp_path, None)
    assert "AI Editor" in xml
    assert "Query to Author" not in xml


def test_the_author_label_is_bold(tmp_path):
    """A bold label is what the eye lands on in a review pane full of grey text."""
    import docx as _docx
    source = tmp_path / "in.docx"
    d = _docx.Document()
    d.add_paragraph("The sample was heated to 40 and held.")
    d.save(str(source))
    out = tmp_path / "author.docx"
    editor.generate_redline_docx(
        str(source), ["The sample was heated to 40 and held."], str(out),
        queries=[{"index": 0, "query": "40 what — °C?", "guard": "x",
                  "audience": "author", "suggestion": None}],
        audience="author")
    with zipfile.ZipFile(str(out)) as z:
        xml = z.read("word/comments.xml").decode("utf8")
    # The label is its own run, and that run is bold.
    head = xml.split("Author:")[0]
    assert "<w:b/>" in head[-400:] or '<w:b ' in head[-400:]
    assert "Author:" in xml and "40 what" in xml


def test_the_internal_copy_has_no_bold_label(tmp_path):
    xml = _redline(tmp_path, None)
    assert "Author:" not in xml
